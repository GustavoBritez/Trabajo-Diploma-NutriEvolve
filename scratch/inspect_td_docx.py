import docx
import json

doc = docx.Document(r'C:\Users\Danie\Desktop\GIT\TD\TD.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

n01_snippets = []
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if any(k in text.lower() for k in ['n01', 'pn1', 'pn2', 'secuencia general', 'turnero', 'seguimiento nutricional', 'actor activo', 'actor pasivo', 'diagrama de secuencia']):
        n01_snippets.append(f"[P{i}] {text}")

with open(r'C:\Users\Danie\Desktop\GIT\TD\scratch\td_docx_n01.txt', 'w', encoding='utf-8') as f:
    f.write("\n\n".join(n01_snippets))

print(f"Found {len(n01_snippets)} matching paragraphs.")
# Also let's print headings
headings = []
for i, p in enumerate(doc.paragraphs):
    if p.style.name.startswith('Heading') or p.text.startswith('N0') or p.text.startswith('G0') or p.text.startswith('T0'):
        headings.append(f"{p.style.name} | {p.text}")

print("=== HEADINGS IN TD.DOCX ===")
for h in headings[:30]:
    print(h)
